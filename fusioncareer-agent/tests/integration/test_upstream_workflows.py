"""Real engine + pinned upstream subprocesses; only external HTTP is simulated."""

import asyncio
import csv
import hashlib
import json
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.algorithms.upstream import OPERATIONS, run_algorithm
from app.config import settings
from app.main import app

VENDOR = Path(__file__).parents[2] / "app/vendor/fusioncareer_algorithm"


@pytest.fixture
def service(tmp_path, monkeypatch):
    calls, uploads = [], []
    state = {"status": 200, "invalid": False, "jobs": []}

    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *args):
            pass

        def send_json(self, value, status=200):
            body = json.dumps(value, ensure_ascii=False).encode()
            self.send_response(status)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def do_GET(self):
            from urllib.parse import parse_qs, urlsplit

            query = parse_qs(urlsplit(self.path).query)
            page = int(query.get("page", [1])[0])
            size = int(query.get("size", [200])[0])
            start = (page - 1) * size
            jobs = state["jobs"]
            self.send_json({"code": 200, "data": {"list": jobs[start:start + size], "total": len(jobs)}})

        def do_POST(self):
            body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            if self.path.endswith("/job-post/batch"):
                uploads.append((self.headers.get("X-Internal-Token"), body))
                self.send_json({"code": 200, "data": {}})
                return
            calls.append(body)
            if state["status"] != 200:
                self.send_json({"error": "unavailable"}, state["status"])
                return
            if "就业顾问" in body["messages"][0]["content"]:
                jobs = json.loads(body["messages"][-1]["content"])["jobs"]
                result = {
                    "order": [job["id"] for job in reversed(jobs)],
                    "reasons": {str(job["id"]): "岗位方向与当前偏好一致。" for job in jobs},
                }
                self.send_json({
                    "id": "test", "object": "chat.completion", "created": 1, "model": "test",
                    "choices": [{"index": 0, "message": {
                        "role": "assistant", "content": json.dumps(result, ensure_ascii=False),
                    }, "finish_reason": "stop"}],
                })
                return
            system_prompt = body["messages"][0]["content"]
            resume = "简历" in system_prompt
            if "岗位目录识别助手" in system_prompt:
                block_id = body["messages"][-1]["content"].split("BLOCK_ID=", 1)[1].splitlines()[0]
                result = {"count": 2, "jobs": [
                    {"companyName": "示例科技有限公司", "positionName": "编辑",
                     "evidenceBlockIds": [block_id]},
                    {"companyName": "示例科技有限公司", "positionName": "记者",
                     "evidenceBlockIds": [block_id]},
                ]}
            elif "TARGET_JOBS_JSON=" in system_prompt:
                targets = json.loads(system_prompt.split("TARGET_JOBS_JSON=", 1)[1])
                result = {"jobs": [{
                    "targetId": target["targetId"],
                    "单位名称": target["companyName"],
                    "岗位名称": target["positionName"],
                    "岗位大类": "企业公司", "岗位二级分类": "民企",
                    "招聘类型": "应届生招聘",
                } for target in targets], "warnings": []}
            elif resume:
                result = {"real_name": "张同学", "gender": 2, "skills": "Python"}
            else:
                result = {"jobs": [{
                    "单位名称": "示例科技有限公司",
                    "岗位名称": "编辑", "岗位大类": "企业公司", "岗位二级分类": "民企",
                    "招聘类型": "小实习", "投递截止日期": "2026-12-31",
                    "其他要求与投递说明": "来源：学院内推",
                }, {
                    "单位名称": "示例科技有限公司", "岗位名称": "记者",
                    "岗位大类": "企业公司", "岗位二级分类": "民企", "招聘类型": "大实习",
                }], "warnings": []}
            self.send_json({
                "id": "test", "object": "chat.completion", "created": 1, "model": "test",
                "choices": [{"index": 0, "message": {
                    "role": "assistant", "content": "invalid" if state["invalid"] else json.dumps(result),
                }, "finish_reason": "stop"}],
                "usage": {"prompt_tokens": 10, "completion_tokens": 20, "total_tokens": 30},
            })

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base = f"http://127.0.0.1:{server.server_port}"
    monkeypatch.setattr(settings, "llm_base_url", base + "/v1")
    monkeypatch.setattr(settings, "llm_api_key", "test-secret")
    monkeypatch.setattr(settings, "llm_model", "test-model")
    monkeypatch.setattr(settings, "backend_base_url", base)
    monkeypatch.setattr(settings, "agent_runtime_dir", str(tmp_path))
    monkeypatch.setattr(settings, "ocr_preload", False)
    monkeypatch.setattr(settings, "internal_service_token", "test-internal")
    monkeypatch.setattr(settings, "agent_admin_token", "test-admin")
    try:
        yield calls, uploads, state, tmp_path
    finally:
        server.shutdown()
        server.server_close()
        thread.join()


def test_vendored_source_matches_pinned_commit():
    manifest = json.loads((VENDOR / "PROVENANCE.json").read_text())
    assert manifest["commit"] == "216e4054acb68ea5b2e7b8b3c82219109912319e"
    for name, digest in manifest["sha256"].items():
        assert hashlib.sha256((VENDOR / name).read_bytes()).hexdigest() == digest, name


def test_latest_recommendation_order_and_reason_are_adapted_without_audit_log(service):
    result = asyncio.run(run_algorithm("job_recommend_rank", {
        "use_llm": True,
        "slots": {"workCities": ["上海"]},
        "jobs": [
            {"id": "101", "companyName": "甲公司", "positionName": "记者",
             "status": "PUBLISHED", "workCity": "上海"},
            {"id": "102", "companyName": "乙公司", "positionName": "编辑",
             "status": "PUBLISHED", "workCity": "上海"},
        ],
    }))
    assert result["method"] == "rules+llm"
    assert [job["id"] for job in result["jobs"]] == ["102", "101"]
    assert all(job["recommendReason"] == "岗位方向与当前偏好一致。" for job in result["jobs"])
    assert not list(service[3].rglob("logs/recommend/*"))


def test_java_job_api_runs_real_upstream_workflow(service, monkeypatch):
    with TestClient(app) as client:
        engine = app.state.workflow_engine
        original = engine.run
        executed = []

        async def tracked(workflow):
            executed.append([node["skill"] for node in workflow["nodes"].values()])
            return await original(workflow)

        monkeypatch.setattr(engine, "run", tracked)
        response = client.post("/api/internal/job/structure", headers={"X-Internal-Token": "test-internal"},
                               json={"text": "学院内推：编辑与记者，12月31日截止"})
        assert response.status_code == 200, response.text
        jobs = response.json()["jobs"]
        assert len(jobs) == 2
        assert jobs[0]["companyName"] == "示例科技有限公司"
        assert jobs[0]["recruitType"] == "SMALL_INTERNSHIP"
        assert jobs[0]["applicationDeadline"].endswith("-12-31")
        assert all(job["status"] == "OFFLINE" for job in jobs)
        assert "native_job_structure" in executed[0]
        assert "学院内推" in service[0][0]["messages"][1]["content"]
        assert app.state.workflow_catalog.get("job_structure")["nodes"]["input"]["inputs"]["json_obj"]["value"]["text"] == ""


@pytest.mark.parametrize("suffix", ["pdf", "docx"])
def test_java_resume_api_runs_download_parse_normalize(service, monkeypatch, tmp_path, suffix):
    from tests.algorithms.test_resume_parser import writePdf

    resume = tmp_path / f"sample.{suffix}"
    if suffix == "pdf":
        writePdf(resume)
    else:
        document = Document()
        document.add_paragraph("张同学 Python")
        document.save(resume)
    ids = []

    async def download(self, user_id, file_id):
        ids.append((user_id, file_id))
        return {"originalName": resume.name}, resume.read_bytes()

    monkeypatch.setattr("app.integrations.backend.BackendClient.read_resume_file", download)
    with TestClient(app) as client:
        response = client.post("/api/internal/resume/parse", headers={"X-Internal-Token": "test-internal"},
                               json={"userId": 11, "fileId": 22})
        assert response.status_code == 200, response.text
        assert response.json()["profilePatch"]["gender"] == "FEMALE"
        assert response.json()["resumePatch"]["skills"] == "Python"
        assert ids == [(11, 22)]
    assert not list(service[3].glob("algorithm/*/resume-*"))


def test_batch_pipeline_exports_and_upload_deduplicates(service):
    with TestClient(app) as client:
        response = client.post("/api/workflows/algorithm_job_pipeline/run", headers={"X-Agent-Admin-Token": "test-admin"},
                               json={"overrides": {"input.json_obj": {
                                   "workspace": "batch", "documents": ["# 招聘岗位汇总\n编辑、记者岗位招聘"],
                               }}})
        assert response.status_code == 200, response.text
        assert response.json()["status"] == "completed", response.text
    root = service[3] / "algorithm/batch"
    assert len(json.loads((root / "data/output/all_positions.json").read_text())) == 2
    from openpyxl import load_workbook
    workbook = load_workbook(root / "data/output/all_positions.xlsx")
    assert workbook.active.max_row == 3
    workbook.close()
    first = asyncio.run(run_algorithm("job_upload", {"workspace": "batch", "dry_run": False}))
    second = asyncio.run(run_algorithm("job_upload", {"workspace": "batch", "dry_run": False}))
    assert first["uploaded"] == 2
    assert second["uploaded"] == 0
    assert second["skipped_dup"] == 2
    assert service[1][0][0] == "test-internal"
    assert all(job["status"] == "OFFLINE" for job in service[1][0][1])


def test_upload_deduplicates_cross_source_campaign_aliases(service):
    root = service[3] / "algorithm/semantic-upload"
    root.mkdir(parents=True)
    shared = {
        "positionName": "行政管培生",
        "jobCategory": "ENTERPRISE",
        "jobSubCategory": "PRIVATE_ENTERPRISE",
        "recruitType": "CAMPUS_RECRUITMENT",
        "reqGradYear": "2027届",
        "reqOther": "网申：https://huicecom.zhiye.com/Campus；邮箱：hr@huice.com",
    }
    rows = [
        {**shared, "sourceUrl": "https://career.fudan.edu.cn/a", "companyName": "慧策集团",
         "workCity": "北京"},
        {**shared, "sourceUrl": "https://jy.xmu.edu.cn/b", "companyName": "北京掌上先机网络科技有限公司",
         "workCity": "北京、天津"},
    ]
    (root / "jobs.json").write_text(json.dumps(rows, ensure_ascii=False), encoding="utf-8")

    result = asyncio.run(run_algorithm("job_upload", {
        "workspace": "semantic-upload", "json": "jobs.json", "dry_run": False,
    }))

    assert result["uploaded"] == 1 and result["skipped_dup"] == 1
    assert len(service[1]) == 1 and len(service[1][0][1]) == 1
    assert set(service[1][0][1][0]["workCities"]) == {"北京", "天津"}


def test_upload_record_dedup_kill_switch_restores_exact_only(service, monkeypatch):
    monkeypatch.setattr(settings, "job_record_dedup_enabled", False)
    root = service[3] / "algorithm/semantic-disabled"
    root.mkdir(parents=True)
    shared = {
        "positionName": "行政管培生",
        "jobCategory": "ENTERPRISE",
        "jobSubCategory": "PRIVATE_ENTERPRISE",
        "recruitType": "CAMPUS_RECRUITMENT",
        "reqGradYear": "2027届",
        "reqOther": "网申：https://huicecom.zhiye.com/Campus；邮箱：hr@huice.com",
    }
    rows = [
        {**shared, "sourceUrl": "https://career.fudan.edu.cn/a", "companyName": "慧策集团"},
        {**shared, "sourceUrl": "https://jy.xmu.edu.cn/b", "companyName": "北京掌上先机网络科技有限公司"},
    ]
    (root / "jobs.json").write_text(json.dumps(rows, ensure_ascii=False), encoding="utf-8")

    result = asyncio.run(run_algorithm("job_upload", {
        "workspace": "semantic-disabled", "json": "jobs.json", "dry_run": False,
    }))

    assert result["uploaded"] == 2 and result["skipped_dup"] == 0


@pytest.mark.parametrize("firstPatch,secondPatch", [
    ({"workCity": "北京"}, {"workCity": "上海"}),
    ({"department": "行政部"}, {"department": "销售部"}),
    ({"recruitType": "CAMPUS_RECRUITMENT"}, {"recruitType": "DAILY_INTERNSHIP"}),
    ({"salaryMin": 8000}, {"salaryMin": 12000}),
])
def test_upload_preserves_semantic_conflicts(service, firstPatch, secondPatch):
    root = service[3] / "algorithm/semantic-conflict"
    root.mkdir(parents=True)
    shared = {
        "companyName": "慧策集团",
        "positionName": "行政管培生",
        "jobCategory": "ENTERPRISE",
        "jobSubCategory": "PRIVATE_ENTERPRISE",
        "recruitType": "CAMPUS_RECRUITMENT",
        "reqGradYear": "2027届",
        "reqOther": "网申：https://huicecom.zhiye.com/Campus；邮箱：hr@huice.com",
    }
    rows = [
        {**shared, **firstPatch, "sourceUrl": "https://career.fudan.edu.cn/a"},
        {**shared, **secondPatch, "sourceUrl": "https://jy.xmu.edu.cn/b"},
    ]
    (root / "jobs.json").write_text(json.dumps(rows, ensure_ascii=False), encoding="utf-8")

    result = asyncio.run(run_algorithm("job_upload", {
        "workspace": "semantic-conflict", "json": "jobs.json", "dry_run": False,
    }))

    assert result["uploaded"] == 2 and result["skipped_dup"] == 0
    assert len(service[1][0][1]) == 2


def test_upload_semantically_matches_existing_backend_job(service):
    root = service[3] / "algorithm/semantic-existing"
    root.mkdir(parents=True)
    existing = {
        "id": "101",
        "sourceUrl": "https://career.fudan.edu.cn/old",
        "companyName": "北京掌上先机网络科技有限公司",
        "positionName": "行政管培生",
        "recruitType": "CAMPUS_RECRUITMENT",
        "reqGradYear": "2027届",
        "workCity": "北京",
        "reqOther": "网申：https://huicecom.zhiye.com/Campus；邮箱：hr@huice.com",
        "status": "PUBLISHED",
        "createdAt": "2026-09-30 20:00:00",
    }
    service[2]["jobs"] = [existing]
    incoming = {
        **existing,
        "id": None,
        "sourceUrl": "https://jy.xmu.edu.cn/new",
        "companyName": "慧策集团",
        "status": "OFFLINE",
    }
    (root / "jobs.json").write_text(json.dumps([incoming], ensure_ascii=False), encoding="utf-8")

    result = asyncio.run(run_algorithm("job_upload", {
        "workspace": "semantic-existing", "json": "jobs.json", "dry_run": False,
    }))

    assert result["uploaded"] == 0 and result["skipped_dup"] == 1
    assert service[1] == []


def test_resume_text_metrics_and_batch_csv(service):
    from tests.algorithms.test_resume_parser import writePdf

    result = asyncio.run(run_algorithm("resume_parse", {"raw_text": "张同学 Python"}))
    assert result["record"]["real_name"] == "张同学"
    assert result["metrics"]["total_tokens"] == 30
    root = service[3] / "algorithm/resumes"
    root.mkdir(parents=True)
    writePdf(root / "one.pdf")
    result = asyncio.run(run_algorithm("resume_batch", {"workspace": "resumes", "files": ["one.pdf"]}))
    with Path(result["file"]).open(encoding="utf-8-sig") as stream:
        assert next(iter(csv.DictReader(stream)))["skills"] == "Python"


def test_upstream_failures_and_paths(service):
    service[2]["invalid"] = True
    with pytest.raises(ValueError, match="invalid JSON"):
        asyncio.run(run_algorithm("job_admin", {"text": "招聘编辑"}))
    service[2]["invalid"] = False
    service[2]["status"] = 402
    with pytest.raises(RuntimeError, match="insufficient balance"):
        asyncio.run(run_algorithm("job_markdown", {"text": "# 新闻记者招聘\n招聘新闻记者"}))
    with pytest.raises(ValueError, match="workspace"):
        asyncio.run(run_algorithm("job_batch", {"workspace": "../escape"}))
    with pytest.raises(ValueError, match="inside its workspace"):
        asyncio.run(run_algorithm("resume_parse", {"file": "../../secret.pdf"}))


def test_all_presets_validate(service):
    with TestClient(app):
        for name in OPERATIONS:
            preset = name if name == "job_structure" else "algorithm_" + name
            workflow = app.state.workflow_catalog.get(preset)
            assert not app.state.workflow_engine.validate(workflow, allow_source_literals_only=True), preset


@pytest.mark.parametrize("suffix,content,status", [(".exe", b"bad", 422), (".pdf", b"x" * (20 * 1024 * 1024 + 1), 413)])
def test_resume_validation_survives_workflow_error_wrapping(service, monkeypatch, suffix, content, status):
    async def download(self, user_id, file_id):
        return {"originalName": "resume" + suffix}, content

    monkeypatch.setattr("app.integrations.backend.BackendClient.read_resume_file", download)
    with TestClient(app) as client:
        response = client.post("/api/internal/resume/parse", headers={"X-Internal-Token": "test-internal"},
                               json={"userId": 1, "fileId": 2})
        assert response.status_code == status
    assert service[0] == []


def test_runtime_preset_is_used_and_failure_not_reported_as_success(service):
    service[2]["invalid"] = True
    with TestClient(app) as client:
        response = client.post("/api/internal/job/structure", headers={"X-Internal-Token": "test-internal"},
                               json={"text": "招聘编辑"})
        assert response.status_code == 400
        preset = app.state.workflow_catalog.get("job_structure")
        preset["nodes"]["result"]["skill"] = "missing_skill"
        response = client.post("/api/internal/job/structure", headers={"X-Internal-Token": "test-internal"},
                               json={"text": "招聘编辑"})
        assert response.status_code == 422


def test_crawler_structuring_uses_same_workflow(service):
    from app.skills.business.crawlers.paths import CrawlPaths
    from app.skills.business.crawlers.store import CrawlStore
    from app.skills.business.crawlers.structure_articles import structureArticles
    from tests.skills.crawlers.test_structure_articles import FakeBackend

    paths = CrawlPaths(service[3] / "wechat")
    store = CrawlStore(paths.database_file)
    store.saveAccount("test", "Test", True)
    markdown = service[3] / "article.md"
    markdown.write_text("# 新闻记者招聘\n投递邮箱：job@example.com")
    store.saveArticle("test", {"title": "新闻记者招聘", "link": "https://example.test/1", "create_time": 1}, markdown, "hash")
    backend = FakeBackend()
    with TestClient(app):
        result = asyncio.run(structureArticles(paths, backend))
    assert result["jobCount"] == 2
    assert "汇总" in service[0][0]["messages"][0]["content"]
    assert all(job["status"] == "OFFLINE" for job in backend.createJobs)


def test_wechat_modes_preserve_history_and_retry_failed_downloads(tmp_path, monkeypatch):
    import time
    from types import SimpleNamespace

    from app.algorithms import upstream_worker as worker

    state = {"articles": [{"title": "岗位一", "link": "https://mp.weixin.qq.com/s/one", "create_time": 1}], "fail": False}

    class Response:
        text = '<div id="js_content"><p>招聘编辑，投递邮箱 job@example.test</p></div>'

        def raise_for_status(self):
            pass

        def json(self):
            return {"base_resp": {"ret": 0}, "publish_page": json.dumps({
                "total_count": len(state["articles"]),
                "publish_list": [{"publish_info": json.dumps({
                    "sent_info": {"time": article["create_time"]},
                    "appmsg_info": [{"title": article["title"], "content_url": article["link"]}],
                })} for article in state["articles"]],
            })}

    def get(url, **kwargs):
        if state["fail"] and "/cgi-bin/" not in url:
            raise RuntimeError("download failed")
        return Response()

    original_load = worker.load_wechat

    def fresh(name, filename):
        fresh_module = original_load(name, filename)
        fresh_module.http_get = get
        fresh_module.time = SimpleNamespace(sleep=lambda _: None, strftime=time.strftime, localtime=time.localtime)
        return fresh_module

    monkeypatch.setattr(worker, "load_wechat", fresh)
    request = {"accounts": [{"fakeid": "test", "name": "测试就业"}]}
    config = {"token": "test", "cookie": "test", "llm_api_key": "test"}
    payload = {"root": str(tmp_path), "operation": "wechat_bootstrap", "request": request, "config": config}
    worker.execute(payload)
    assert list(tmp_path.glob("data/articles/**/*.md"))
    original_history = (tmp_path / "history.json").read_bytes()
    original_log = (tmp_path / "wx_poc.txt").read_bytes()
    state["articles"].insert(0, {"title": "岗位二", "link": "https://mp.weixin.qq.com/s/two", "create_time": 2})
    state["fail"] = True
    payload["operation"] = "wechat_daily"
    with pytest.raises(RuntimeError, match="download failed"):
        worker.execute(payload)
    assert (tmp_path / "history.json").read_bytes() == original_history
    assert (tmp_path / "wx_poc.txt").read_bytes() == original_log
    state["fail"] = False
    report = worker.execute(payload)
    assert report["total_new"] == 1
    assert (tmp_path / "daily_report.jsonl").exists()
    assert Path(report["daily_directory"], "新增统计.md").exists()
    payload["operation"] = "wechat_archive"
    assert worker.execute(payload)["accounts"] == 1


def test_session_sync_calls_upstream_and_returns_no_credentials(tmp_path, monkeypatch):
    from types import SimpleNamespace

    from app.algorithms import upstream_worker as worker

    module = SimpleNamespace()

    def sync(url):
        assert url == "http://127.0.0.1:9222"
        module.CONFIG_FILE.write_text(json.dumps({"token": "private", "cookie": "private"}))
        return True

    module.run_cdp = sync
    monkeypatch.setattr(worker, "load_wechat", lambda *args: module)
    result = worker.wechat("wechat_sync_session", {"cdp_url": "http://127.0.0.1:9222"}, tmp_path, {})
    assert result == {"synced": True}
    assert module.CONFIG_FILE.stat().st_mode & 0o777 == 0o600


def test_timeout_terminates_worker_waiting_for_workspace_lock(service, monkeypatch):
    import fcntl

    root = service[3] / "algorithm/locked"
    root.mkdir(parents=True)
    processes = []
    create_process = asyncio.create_subprocess_exec

    async def tracked(*args, **kwargs):
        process = await create_process(*args, **kwargs)
        processes.append(process)
        return process

    monkeypatch.setattr(asyncio, "create_subprocess_exec", tracked)
    monkeypatch.setattr(settings, "algorithm_timeout_seconds", 0.1)
    with (root / ".lock").open("a") as lock:
        fcntl.flock(lock, fcntl.LOCK_EX)
        with pytest.raises(TimeoutError):
            asyncio.run(run_algorithm("job_batch", {"workspace": "locked"}))
    assert processes[0].returncode is not None


def test_parallel_batches_share_dedup_state_without_duplicate_writes(service):
    async def run():
        request = {"workspace": "parallel", "documents": ["# 招聘编辑\n编辑和记者招聘"]}
        return await asyncio.gather(run_algorithm("job_batch", request), run_algorithm("job_batch", request))

    results = asyncio.run(run())
    assert all(len(result["positions"]) == 2 for result in results)
    with (service[3] / "algorithm/parallel/data/output/all_positions.csv").open(encoding="utf-8-sig") as stream:
        assert len(list(csv.DictReader(stream))) == 2
