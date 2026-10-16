"""Real file download HTTP and pinned parser transport; all documents are synthetic."""
import asyncio
import json
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.algorithms.upstream import run_algorithm
from app.chat.tools import runTool
from app.config import settings
from app.integrations.backend import BackendClient
from app.main import app
from tests.algorithms.test_resume_advice import context, report
from tests.algorithms.test_resume_parser import writePdf


@pytest.fixture
def transport(tmp_path, monkeypatch):
    state = {"filename": "resume.pdf", "bytes": b"", "requests": [],
             "mode": "ok", "model_calls": 0, "get_count": 0, "context_calls": []}

    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *args):
            pass

        def send(self, value, status=200, content_type="application/json"):
            body = value if isinstance(value, bytes) else json.dumps(value).encode()
            self.send_response(status)
            self.send_header("Content-Type", content_type)
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            try:
                self.wfile.write(body)
            except BrokenPipeError:
                pass

        def do_GET(self):
            state["get_count"] += 1
            assert self.headers.get("X-Internal-Token") == "test-internal"
            if self.path == "/internal/resume-file/11/list":
                self.send({"code": 200, "data": [{"id": "22", "originalName": state["filename"]}]})
            elif self.path == "/internal/resume-file/22/download":
                self.send(state["bytes"], content_type="application/octet-stream")
            else:
                self.send({"code": 404}, status=404)

        def do_POST(self):
            request = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            if self.path.startswith("/internal/agent/tools/"):
                assert self.headers.get("X-Internal-Token") == "test-internal"
                assert self.headers.get("X-Agent-Context") == "signed-test-context"
                time.sleep(0.1)
                if self.path.endswith("/resume_advice_context"):
                    state["context_calls"].append(request)
                    data = {**context(), "resumeSource": "file", "resumeVersion": "file:22:v1"}
                else:
                    data = {"id": "123", "positionName": "Reporter"}
                self.send({"code": 200, "data": data})
                return
            state["model_calls"] += 1
            state["requests"].append(request)
            if state["mode"] == "http_error":
                self.send({"error": {"message": "temporary", "type": "server_error"}}, status=503)
                return
            if state["mode"] == "slow":
                time.sleep(0.3)
            invalid = state["mode"] == "json_retry" and state["model_calls"] == 1
            content = "invalid" if invalid else json.dumps({"gender": 2, "skills": "Python"})
            if state["mode"] == "advice":
                content = json.dumps({**report(), "resumeVersion": "file:22:v1"})
            self.send({"id": "test", "object": "chat.completion", "created": 1,
                       "model": "deepseek-v4-flash", "choices": [{"index": 0,
                       "message": {"role": "assistant", "content": content},
                       "finish_reason": "length" if state["mode"] == "truncated" else "stop"}],
                       "usage": {"prompt_tokens": 20, "completion_tokens": 12, "total_tokens": 32}})

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base = f"http://127.0.0.1:{server.server_port}"
    for key, value in {"llm_base_url": base + "/v1", "llm_api_key": "synthetic-key",
                       "llm_model": "deepseek-v4-flash", "backend_base_url": base,
                       "agent_runtime_dir": str(tmp_path / "runtime"), "ocr_preload": False,
                       "internal_service_token": "test-internal", "agent_admin_token": "test-admin"}.items():
        monkeypatch.setattr(settings, key, value)
    monkeypatch.setattr("app.main.backend_client", BackendClient())
    try:
        yield state
    finally:
        server.shutdown()
        server.server_close()
        thread.join()


@pytest.mark.parametrize("suffix", ["pdf", "docx"])
def test_real_http_download_and_pinned_parse_share_bounded_model_config(transport, tmp_path, suffix):
    path = tmp_path / f"resume.{suffix}"
    if suffix == "pdf":
        writePdf(path)
    else:
        document = Document()
        document.add_paragraph("Synthetic student. Python. Reporting project.")
        document.save(path)
    transport.update(filename=path.name, bytes=path.read_bytes())
    with TestClient(app) as client:
        response = client.post("/api/internal/resume/parse", headers={"X-Internal-Token": "test-internal"},
                               json={"userId": 11, "fileId": 22})
    assert response.status_code == 200, response.text
    assert response.json()["resumePatch"]["skills"] == "Python"
    assert transport["get_count"] == 2
    assert transport["model_calls"] == 1
    request = transport["requests"][0]
    assert request["model"] == "deepseek-v4-flash"
    assert request["max_tokens"] == 8192
    assert request["thinking"] == {"type": "disabled"}
    assert request["response_format"] == {"type": "json_object"}
    assert request["temperature"] == 0


@pytest.mark.parametrize("mode, expected_calls, expected_status", [
    ("json_retry", 2, 200), ("http_error", 1, 500), ("truncated", 1, 500),
])
def test_json_retry_preserved_but_sdk_network_retry_disabled(transport, tmp_path, monkeypatch,
                                                           mode, expected_calls, expected_status):
    path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("Synthetic student Python")
    document.save(path)
    transport.update(filename=path.name, bytes=path.read_bytes(), mode=mode)
    with TestClient(app, raise_server_exceptions=False) as client:
        response = client.post("/api/internal/resume/parse", headers={"X-Internal-Token": "test-internal"},
                               json={"userId": 11, "fileId": 22})
    assert response.status_code == expected_status
    assert transport["model_calls"] == expected_calls


def test_http_parse_budget_covers_download_and_cancels_workflow(transport, monkeypatch):
    cancelled = []
    async def slow_workflow(*args):
        try:
            await asyncio.sleep(1)
        finally:
            cancelled.append(True)
    monkeypatch.setattr(settings, "resume_parse_timeout_seconds", 0.02)
    monkeypatch.setattr("app.api.routers.internal.runAlgorithmWorkflow", slow_workflow)
    with TestClient(app) as client:
        response = client.post("/api/internal/resume/parse", headers={"X-Internal-Token": "test-internal"},
                               json={"userId": 11, "fileId": 22})
    assert response.status_code == 504
    assert cancelled == [True]
    assert settings.algorithm_timeout_seconds == 1800


def test_file_advice_helper_uses_parse_budget_with_real_backend_http_and_model(transport, monkeypatch):
    transport["mode"] = "advice"
    monkeypatch.setattr(settings, "ai_chat_read_tool_timeout_seconds", 0.02)
    monkeypatch.setattr(settings, "ai_chat_parse_timeout_seconds", 0.5)
    monkeypatch.setattr(settings, "ai_chat_tool_timeouts", {})
    assert settings.tool_timeout("resume_advice_context") == 0.5
    assert settings.tool_timeout("get_job") == 0.02

    async def advise_and_query():
        backend = BackendClient()
        try:
            result = await runTool(backend, "signed-test-context", "advise_resume",
                                   '{"source":"file","fileId":"22"}', "file-advice")
            assert result.succeeded is True
            assert result.proposed_action is None
            assert result.presentation["type"] == "resume_advice"
            rendered = result.presentation["report"]
            assert rendered["status"] == "ok"
            assert rendered["resumeSource"] == "file"
            assert rendered["requirements"][0]["evidenceIds"] == ["resume-internship"]
            with pytest.raises(httpx.ReadTimeout):
                await backend.run_agent_tool("get_job", {"jobId": "123"}, "signed-test-context")
        finally:
            await backend.close()

    asyncio.run(advise_and_query())
    assert transport["context_calls"] == [{"source": "file", "fileId": "22"}]
    assert transport["model_calls"] == 1


def test_sdk_timeout_does_not_retry_network_request(transport):
    transport["mode"] = "slow"
    with pytest.raises(RuntimeError, match="timed out"):
        asyncio.run(run_algorithm("resume_parse", {"raw_text": "Synthetic student Python",
                                                  "config": {"llm_timeout_seconds": 0.1}}))
    assert transport["model_calls"] == 1


@pytest.mark.parametrize("clock, expected_calls, expected_timeouts", [
    ([10, 10, 45], 2, [60, 25]), ([10, 10, 71], 1, [60]),
])
def test_json_repair_uses_remaining_shared_model_budget(monkeypatch, tmp_path,
                                                      clock, expected_calls, expected_timeouts):
    from importlib import import_module
    from types import SimpleNamespace

    from app.algorithms import upstream_worker

    deepseek_client = import_module("resume_parser.llm.deepseek_client")
    calls, retry_limits = [], []
    class FakeOpenAI:
        def __init__(self, **kwargs):
            self.chat = SimpleNamespace(completions=SimpleNamespace(create=self.create))
        def with_options(self, **kwargs):
            retry_limits.append(kwargs["max_retries"])
            return self
        def create(self, **kwargs):
            calls.append(kwargs)
            content = "invalid" if len(calls) == 1 else '{"skills":"Python"}'
            return SimpleNamespace(choices=[SimpleNamespace(finish_reason="stop", message=SimpleNamespace(
                content=content))], usage=None)

    monkeypatch.setattr(deepseek_client, "OpenAI", FakeOpenAI)
    readings = iter(clock)
    monkeypatch.setattr(upstream_worker.time, "monotonic", lambda: next(readings))
    cfg = {"llm_api_key": "test-key", "llm_model": "deepseek-v4-flash", "llm_base_url": "http://local",
           "llm_timeout_seconds": 60, "llm_max_tokens": 8192,
           "llm_extra_body": {"thinking": {"type": "disabled"}}}
    if expected_calls == 1:
        with pytest.raises(RuntimeError, match="deadline exceeded"):
            upstream_worker.resumes("resume_parse", {"raw_text": "Synthetic Python"}, tmp_path, cfg)
    else:
        result = upstream_worker.resumes("resume_parse", {"raw_text": "Synthetic Python"}, tmp_path, cfg)
        assert result["record"]["skills"] == "Python"
    assert len(calls) == expected_calls
    assert [call["timeout"] for call in calls] == expected_timeouts
    assert retry_limits == [0]
    assert all(call["max_tokens"] == 8192 for call in calls)
    assert all(call["extra_body"] == {"thinking": {"type": "disabled"}} for call in calls)


def test_http_deadline_kills_actual_worker_waiting_for_workspace_lock(transport, tmp_path, monkeypatch):
    import fcntl
    from types import SimpleNamespace

    from app.algorithms import upstream

    document = Document()
    document.add_paragraph("Synthetic Python")
    path = tmp_path / "resume.docx"
    document.save(path)
    transport.update(filename=path.name, bytes=path.read_bytes())
    root = tmp_path / "runtime/algorithm/run-locktest"
    root.mkdir(parents=True)
    processes = []
    original = asyncio.create_subprocess_exec
    async def tracked_process(*args, **kwargs):
        process = await original(*args, **kwargs)
        processes.append(process)
        return process
    monkeypatch.setattr(upstream.uuid, "uuid4", lambda: SimpleNamespace(hex="locktest"))
    monkeypatch.setattr(upstream.asyncio, "create_subprocess_exec", tracked_process)
    monkeypatch.setattr(settings, "resume_parse_timeout_seconds", 0.5)
    with (root / ".lock").open("a") as lock:
        fcntl.flock(lock, fcntl.LOCK_EX)
        with TestClient(app) as client:
            response = client.post("/api/internal/resume/parse", headers={"X-Internal-Token": "test-internal"},
                                   json={"userId": 11, "fileId": 22})
        fcntl.flock(lock, fcntl.LOCK_UN)
    assert response.status_code == 504
    assert len(processes) == 1
    assert processes[0].returncode is not None
    assert transport["model_calls"] == 0
    assert settings.algorithm_timeout_seconds == 1800
